from __future__ import annotations

import asyncio
import io
from datetime import datetime
from types import SimpleNamespace

import openpyxl


def test_der_loader_sheet_with_field_value_header_parses_as_tall_updates() -> None:
    from services.spreadsheet_parser import parse_spreadsheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Case Update"
    ws.append(["DER No", "Field Name", "Field Value"])
    ws.append([26162819616, "Letter template", "Approved"])
    ws.append([26170559280, "Letter template", "Duplicate Request"])

    buf = io.BytesIO()
    wb.save(buf)

    module, rows = parse_spreadsheet(buf.getvalue())

    assert module == "der"
    assert [(row.module, row.case_id, row.field_name, row.new_value) for row in rows] == [
        ("der", "26162819616", "Letter template", "Approved"),
        ("der", "26170559280", "Letter template", "Duplicate Request"),
    ]


def test_exrm_loader_sheet_with_field_value_header_routes_numeric_cases_to_ire_mgmt() -> None:
    from services.spreadsheet_parser import parse_spreadsheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Case Update"
    ws.append(["EXRM No", "Field Name", "Field Value"])
    ws.append([25223289138, "Status", "Closed"])
    ws.append([25223289138, "Status Reason", "Invalid case"])

    buf = io.BytesIO()
    wb.save(buf)

    module, rows = parse_spreadsheet(buf.getvalue())

    assert module == "ire_mgmt"
    assert [(row.module, row.case_id, row.field_name, row.new_value) for row in rows] == [
        ("ire_mgmt", "25223289138", "Status", "Closed"),
        ("ire_mgmt", "25223289138", "Status Reason", "Invalid case"),
    ]


def test_appeals_loader_sheet_with_appeal_no_header_parses_as_tall_updates() -> None:
    from services.spreadsheet_parser import parse_spreadsheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Case Update"
    ws.append(["Appeal No", "Field Name", "Field Value", "Reason For Change"])
    ws.append(["A26123456789", "Status", "Closed", "Post-closure correction"])

    buf = io.BytesIO()
    wb.save(buf)

    module, rows = parse_spreadsheet(buf.getvalue())

    assert module == "appeals_mgmt"
    assert [(row.module, row.case_id, row.field_name, row.new_value, row.reason) for row in rows] == [
        ("appeals_mgmt", "A26123456789", "Status", "Closed", "Post-closure correction"),
    ]


def test_medical_appeals_backend_template_normalizes_field_labels() -> None:
    from services.field_mapper import FieldMapper
    from services.spreadsheet_parser import parse_spreadsheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Back end update multiple field "
    ws.append(["Module", "Medical Appeals"])
    ws.append([])
    ws.append(["Enter one field upate per row"])
    ws.append([
        "MHK Auth/Referral/Case number",
        "Field Name as it appears in the UI",
        "Current Value",
        "New Value",
        "Reason for update",
    ])
    ws.append(["A2618278831", "Decision maker ", "Christopher Bender", "blank", "data entry error"])
    ws.append([
        "A2618370684",
        "Provider Oral Notification (Date)",
        datetime(2026, 7, 4, 8, 51),
        "Blank",
        "remove provider oral notification",
    ])

    buf = io.BytesIO()
    wb.save(buf)

    mapper = FieldMapper()
    mapper.load_dir("meridian/backend/mappings")
    module, rows = parse_spreadsheet(buf.getvalue(), mapper=mapper)

    assert module == "appeals_mgmt"
    assert [(row.case_id, row.field_name, row.new_value) for row in rows] == [
        ("A2618278831", "Appeals Decision Maker", None),
        ("A2618370684", "Provider Oral Notification Date", None),
    ]


def test_wide_parser_does_not_use_request_type_old_for_original_request_reason_context() -> None:
    from services.spreadsheet_parser import parse_spreadsheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["DER"])
    ws.append(["Case #", "Request Type (OLD)", "Original Request Reason (NEW)"])
    ws.append(["26162457864", "Coverage Determination", "Prior Authorization"])

    buf = io.BytesIO()
    wb.save(buf)

    module, rows = parse_spreadsheet(buf.getvalue())

    assert module == "der"
    assert [(row.field_name, row.new_value, row.old_value) for row in rows] == [
        ("Original Request Reason", "Prior Authorization", None),
    ]


def test_case_management_bee_current_new_value_sheet_parses_as_tall_with_document_context() -> None:
    from services.spreadsheet_parser import parse_spreadsheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BEE 06-18-2026"
    ws.append([None, None, None, "Case Management"])
    ws.append([])
    ws.append(["Enter one field update per row"])
    ws.append([
        "Member ID",
        "Internal ID",
        "MHK Auth/Referral/Case number",
        "Field Name as it appears in the UI",
        "Assessment Name",
        '"Created at" date',
        "Current Value",
        "New Value",
    ])
    ws.append([
        36237989,
        "N/A",
        "P26149555568359",
        "Created At",
        "HSC_11434_USPS_Tracking____USPS_TEST1.pdf",
        datetime(2026, 6, 9, 9, 25),
        datetime(2026, 6, 9, 9, 25),
        datetime(2026, 6, 8, 9, 26),
    ])

    buf = io.BytesIO()
    wb.save(buf)

    module, rows = parse_spreadsheet(buf.getvalue())

    assert module == "member_assessment"
    assert len(rows) == 1
    row = rows[0]
    assert row.case_id == "P26149555568359"
    assert row.field_name == "Created At"
    assert row.assessment_name == "HSC_11434_USPS_Tracking____USPS_TEST1.pdf"
    assert row.created_at == "2026-06-09 09:25:00"
    assert row.semantic_context == {
        "assessment_name": "HSC_11434_USPS_Tracking____USPS_TEST1.pdf",
        "document_name": "HSC_11434_USPS_Tracking____USPS_TEST1.pdf",
        "source_created_at": "2026-06-09 09:25:00",
    }
    assert row.old_value == "2026-06-09 09:25:00"
    assert row.new_value == "2026-06-08 09:26:00"


def test_cigna_case_management_bee_guidance_row_parses_without_llm_fallback() -> None:
    from services.field_mapper import FieldMapper
    from services.spreadsheet_parser import parse_spreadsheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BEE "
    ws.append(["Module", None, None, "Case Management"])
    ws.append([])
    ws.append(["Enter one field upate per row"])
    ws.append([
        "Member ID",
        "Internal ID",
        "MHK Auth/Referral/Case number",
        "Field Name as it appears in the UI",
        "Assessment Name",
        '"Created at" date ',
        "Current Value",
        "New Value",
        "Reason for update",
    ])
    ws.append([
        "Member ID",
        "Assessment ID ",
        "Auth Reference/Review #",
        "Name of the field: Status",
        "Must read identical to MHK",
        "Must Match MHK",
        "Status",
        "Denied ",
        "The reason the update is needed",
    ])
    ws.append([
        "U6663571501",
        8400236,
        "P21118795691588",
        "Complete",
        "Intervention Assessment",
        datetime(2026, 7, 8, 0, 5),
        "Complete",
        "Cancelled",
        "CM added intervention assessment to wrong chart",
    ])
    ws.append([
        36237989,
        "N/A",
        "P26149191811482",
        "Created At",
        "HSC_11435_USPS_TRACKING_LEACOCK.pdf",
        datetime(2026, 6, 9, 10, 3),
        datetime(2026, 6, 9, 10, 3),
        datetime(2026, 6, 8, 10, 3),
        "Dates need to match",
    ])

    buf = io.BytesIO()
    wb.save(buf)

    mapper = FieldMapper()
    mapper.load_dir("meridian/backend/mappings")
    module, rows = parse_spreadsheet(buf.getvalue(), mapper=mapper)

    assert module == "member_assessment"
    assert [(row.case_id, row.field_name, row.old_value, row.new_value) for row in rows] == [
        ("P21118795691588", "Status", "Complete", "Cancelled"),
        ("P26149191811482", "Created At", "2026-06-09 10:03:00", "2026-06-08 10:03:00"),
    ]
    assert rows[0].semantic_context == {
        "assessment_name": "Intervention Assessment",
        "document_name": "Intervention Assessment",
        "source_created_at": "2026-07-08 00:05:00",
    }
    assert rows[1].semantic_context == {
        "assessment_name": "HSC_11435_USPS_TRACKING_LEACOCK.pdf",
        "document_name": "HSC_11435_USPS_TRACKING_LEACOCK.pdf",
        "source_created_at": "2026-06-09 10:03:00",
    }


def test_authorization_becu_wide_sheet_uses_jira_context_and_workflow_flag() -> None:
    from services.field_mapper import FieldMapper
    from services.spreadsheet_parser import infer_module_from_case_ids, parse_spreadsheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Case Updates"
    ws.append([
        "Member ID (Optional)",
        "Case # (Mandatory)",
        "Revision #",
        "Field A (Old Value)",
        "Field A (New Value)",
        "Field B (Old Value)",
        "Field B (New Value)",
        "Remove from Workflow Y/N",
    ])
    ws.append([None, None, None, "STATUS", "STATUS", "STATUS REASON", "STATUS REASON", None])
    ws.append([
        "001584113",
        "0010389202",
        None,
        "In progress",
        "Complete",
        "clinical review",
        "Decisioned",
        "Y",
    ])
    ws.append([
        "002314650",
        "PO01010182847",
        None,
        "In progress",
        "Complete",
        "clinical review",
        "Decisioned",
        "Y",
    ])

    buf = io.BytesIO()
    wb.save(buf)

    mapper = FieldMapper()
    mapper.load_dir("meridian/backend/mappings")

    assert infer_module_from_case_ids(["PO01010182847"]) is None

    module, rows = parse_spreadsheet(
        buf.getvalue(),
        mapper=mapper,
        context_text="Please have the authorizations listed in the attachment force closed.",
    )

    assert module == "authorizations"
    assert [
        (row.module, row.case_id, row.field_name, row.old_value, row.new_value)
        for row in rows
    ] == [
        ("authorizations", "0010389202", "Event Status", "In progress", "Complete"),
        ("authorizations", "0010389202", "Status Reason", "clinical review", "Decisioned"),
        ("authorizations", "0010389202", "Terminate Workflow", None, "Y"),
        ("authorizations", "PO01010182847", "Event Status", "In progress", "Complete"),
        ("authorizations", "PO01010182847", "Status Reason", "clinical review", "Decisioned"),
        ("authorizations", "PO01010182847", "Terminate Workflow", None, "Y"),
    ]


def test_calc_expression_override_uses_pending_request_type_placeholder() -> None:
    from models.schemas import ChangeRow
    from routers.generate import _select_calc_expression
    from services.case_lookup import resolve_via_calc_expression

    class Cursor:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def execute(self, sql):
            self.pool.sql = sql

        async def fetchall(self):
            return [(18,)]

    class Conn:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        def cursor(self):
            return Cursor(self.pool)

    class Pool:
        sql = ""

        def acquire(self):
            return Conn(self)

    change = ChangeRow(
        module="der",
        case_id="26162457864",
        field_name="RequestReason",
        column_name="cdtype_id",
        new_value="Prior Authorization",
    )
    overrides = {
        "der": {
            "cdtype_id": {
                "required_batch_values": ["request_type"],
                "calc_expression": (
                    "SELECT id FROM cdtype WHERE description = '#RequestReason#' "
                    "AND request_type_id = #request_type#"
                ),
            }
        }
    }
    calc = _select_calc_expression(
        "SELECT id FROM cdtype WHERE description = '#RequestReason#'",
        change,
        {"request_type": "4"},
        overrides,
    )

    pool = Pool()
    result = asyncio.run(
        resolve_via_calc_expression(
            pool,
            calc,
            "Prior Authorization",
            "RequestReason",
            "26162457864",
            {"request_type": "4"},
        )
    )

    assert result == "18"
    assert "request_type_id = 4" in pool.sql
    assert "SELECT request_type FROM der" not in pool.sql


def test_calc_expression_distinguishes_display_and_column_placeholder_case() -> None:
    from services.case_lookup import resolve_via_calc_expression

    class Cursor:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def execute(self, sql):
            self.pool.sql = sql

        async def fetchall(self):
            return [(170,)]

    class Conn:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        def cursor(self):
            return Cursor(self.pool)

    class Pool:
        sql = ""

        def acquire(self):
            return Conn(self)

    pool = Pool()
    result = asyncio.run(
        resolve_via_calc_expression(
            pool,
            (
                "SELECT id FROM der_status_reason "
                "WHERE description = '#Status Reason#' "
                "AND der_status_id = #status# "
                "AND display_message = '#Status#'"
            ),
            "Medical Necessity",
            "Status Reason",
            "26148341125",
            {
                "status": "15",
                "Status": "Overturned",
                "Status Reason": "Medical Necessity",
            },
        )
    )

    assert result == "170"
    assert "description = 'Medical Necessity'" in pool.sql
    assert "der_status_id = 15" in pool.sql
    assert "display_message = 'Overturned'" in pool.sql


def test_calc_expression_override_scopes_ire_status_reason_by_current_status() -> None:
    from models.schemas import ChangeRow
    from routers.generate import _load_calc_expression_overrides, _select_calc_expression
    from services.case_lookup import resolve_via_calc_expression

    class Cursor:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def execute(self, sql):
            self.pool.sql = sql

        async def fetchall(self):
            return [(10,)]

    class Conn:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        def cursor(self):
            return Cursor(self.pool)

    class Pool:
        sql = ""

        def acquire(self):
            return Conn(self)

    change = ChangeRow(
        module="ire_mgmt",
        case_id="26110755711",
        field_name="Status Reason",
        column_name="ire_status_reason_id",
        new_value="Invalid case",
    )
    calc = _select_calc_expression(
        "SELECT id FROM ire_status_reason WHERE description = '#Status Reason#'",
        change,
        {},
        _load_calc_expression_overrides(),
    )

    pool = Pool()
    result = asyncio.run(
        resolve_via_calc_expression(
            pool,
            calc,
            "Invalid case",
            "Status Reason",
            "26110755711",
            {},
        )
    )

    assert result == "10"
    assert "FROM ire_status_reason" in pool.sql
    assert "ire_no = '26110755711'" in pool.sql


def test_calc_expression_override_scopes_ire_status_reason_by_pending_status() -> None:
    from models.schemas import ChangeRow
    from routers.generate import _load_calc_expression_overrides, _select_calc_expression
    from services.case_lookup import resolve_via_calc_expression

    class Cursor:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def execute(self, sql):
            self.pool.sql = sql

        async def fetchall(self):
            return [(10,)]

    class Conn:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        def cursor(self):
            return Cursor(self.pool)

    class Pool:
        sql = ""

        def acquire(self):
            return Conn(self)

    change = ChangeRow(
        module="ire_mgmt",
        case_id="25223289138",
        field_name="Status Reason",
        column_name="ire_status_reason_id",
        new_value="Invalid case",
    )
    calc = _select_calc_expression(
        "SELECT id FROM ire_status_reason WHERE description = '#Status Reason#'",
        change,
        {"ire_status_id": "3"},
        _load_calc_expression_overrides(),
    )

    pool = Pool()
    result = asyncio.run(
        resolve_via_calc_expression(
            pool,
            calc,
            "Invalid case",
            "Status Reason",
            "25223289138",
            {"ire_status_id": "3"},
        )
    )

    assert result == "10"
    assert "FROM ire_status_reason" in pool.sql
    assert "ire_status_id = 3" in pool.sql
    assert "FROM ire_mgmt" not in pool.sql


def test_calc_expression_override_scopes_ire_status_by_request_type() -> None:
    from models.schemas import ChangeRow
    from routers.generate import _load_calc_expression_overrides, _select_calc_expression
    from services.case_lookup import resolve_via_calc_expression

    class Cursor:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def execute(self, sql):
            self.pool.sql = sql

        async def fetchall(self):
            return [(53,)]

    class Conn:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        def cursor(self):
            return Cursor(self.pool)

    class Pool:
        sql = ""

        def acquire(self):
            return Conn(self)

    change = ChangeRow(
        module="ire_mgmt",
        case_id="25223289138",
        field_name="Status",
        column_name="ire_status_id",
        new_value="Closed",
    )
    calc = _select_calc_expression(
        None,
        change,
        {},
        _load_calc_expression_overrides(),
    )

    pool = Pool()
    result = asyncio.run(
        resolve_via_calc_expression(
            pool,
            calc,
            "Closed",
            "Status",
            "25223289138",
            {},
        )
    )

    assert result == "53"
    assert "FROM ire_status" in pool.sql
    assert "ire_request_type_id" in pool.sql
    assert "ire_no = '25223289138'" in pool.sql


def test_reference_or_calc_fields_require_id_resolution_before_noop_check() -> None:
    from routers.generate import _requires_value_resolution

    plain = SimpleNamespace(
        calc_expression=None,
        lov_expression=None,
        parent_columns=None,
        ref_table=None,
        ref_id_col=None,
        ref_desc_col=None,
    )
    ref = SimpleNamespace(
        calc_expression=None,
        lov_expression=None,
        parent_columns=None,
        ref_table="note_types",
        ref_id_col="id",
        ref_desc_col="description",
    )
    calc = SimpleNamespace(
        calc_expression="SELECT id FROM note_types WHERE description = '#Letter template#'",
        lov_expression=None,
        parent_columns=None,
        ref_table="note_types",
        ref_id_col="id",
        ref_desc_col="description",
    )

    assert _requires_value_resolution(plain) is False
    assert _requires_value_resolution(ref) is True
    assert _requires_value_resolution(calc) is True


def test_bulk_preload_does_not_mark_ref_table_display_match_already_done(monkeypatch) -> None:
    from models.schemas import ChangeRow
    from routers import generate

    async def fake_bulk_find_record_ids(pool, module, case_ids):
        return {"25223289138": 32316}

    async def fake_bulk_get_current_values(pool, table, pk_col, record_ids, columns):
        return {(32316, "ire_status_id"): ("3", "Closed")}

    class Mapper:
        def find_for_module(self, module, field_name):
            return SimpleNamespace(
                calc_expression=None,
                lov_expression=None,
                parent_columns=None,
                ref_table="ire_status",
                ref_id_col="id",
                ref_desc_col="description",
            )

    monkeypatch.setattr(generate, "bulk_find_record_ids", fake_bulk_find_record_ids)
    monkeypatch.setattr(generate, "bulk_get_current_values", fake_bulk_get_current_values)

    change = ChangeRow(
        module="ire_mgmt",
        case_id="25223289138",
        field_name="Status",
        column_name="ire_status_id",
        new_value="Closed",
    )

    asyncio.run(generate._bulk_preload([change], object(), Mapper()))

    assert change.record_id == 32316
    assert change.current_db_value == "Closed"
    assert change.already_done is False


def test_calc_expression_prefers_most_specific_note_type_context_candidate() -> None:
    from services.case_lookup import resolve_via_calc_expression

    class Cursor:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def execute(self, sql):
            self.pool.sql_calls.append(sql)

        async def fetchall(self):
            if len(self.pool.sql_calls) == 1:
                return [(4, "APPROVED"), (2207, "Approved")]
            return [
                (4, "APPROVED", 3371, 0),
                (2207, "Approved", 18285, 4),
            ]

    class Conn:
        def __init__(self, pool):
            self.pool = pool

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        def cursor(self):
            return Cursor(self.pool)

    class Pool:
        def __init__(self):
            self.sql_calls = []

        def acquire(self):
            return Conn(self)

    calc = """
        SELECT DISTINCT nt.ID, nt.description
        FROM note_types nt
        INNER JOIN note_type_context ntc ON ntc.note_type_id = nt.ID
        INNER JOIN der d ON d.der_no = '#DER No#'
        WHERE nt.description = '#Letter template#'
    """

    pool = Pool()
    result = asyncio.run(
        resolve_via_calc_expression(
            pool,
            calc,
            "Approved",
            "Letter template",
            "26162819616",
            record_id=29742661,
        )
    )

    assert result == "2207"
    assert len(pool.sql_calls) == 2
    assert "specificity" in pool.sql_calls[1]
    assert "`ntc`.`der_request_type_id`" in pool.sql_calls[1]
