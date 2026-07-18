from __future__ import annotations

import unittest
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from services.spreadsheet_parser import extract_case_ids, infer_module_from_case_ids, parse_spreadsheet
from services.text_extra_fields import extract_deterministic_ticket_extra_fields


@dataclass
class _Mapping:
    ui_name: str


class _Mapper:
    def __init__(self) -> None:
        self._fields = {
            ("provider_appeals_mgmt", "Category"): _Mapping("Category"),
            ("provider_appeals_mgmt", "Sub Category"): _Mapping("Sub Category"),
            ("provider_appeals_mgmt", "Status"): _Mapping("Status"),
            ("provider_appeals_mgmt", "Status Reason"): _Mapping("Status Reason"),
            ("provider_appeals_mgmt", "Case Closure Date"): _Mapping("Case Closure Date"),
            ("provider_appeals_mgmt", "Dispute Request Reason"): _Mapping("Dispute Request Reason"),
            ("provider_appeals_mgmt", "Terminate Workflow"): _Mapping("Terminate Workflow"),
            ("appeals_mgmt", "Appeals Category"): _Mapping("Appeals Category"),
            ("der", "Terminate Workflow"): _Mapping("Terminate Workflow"),
        }

    def find_for_module(self, module: str, ui_name: str):
        return self._fields.get((module, ui_name))

    def get_by_ui(self, ui_name: str):
        return [field for (module, name), field in self._fields.items() if name == ui_name]


class TextExtraFieldsTest(unittest.TestCase):
    def test_extracts_provider_appeals_ticket_level_fields(self) -> None:
        text = """
        Back end case update:
        Category - Medicare Coverage Determination
        Subcategory - Local Coverage Determination
        Status - Overturned
        Status Reason - Claims Processed Incorrectly
        Closed Date - 6/4/2026
        Case notes - Closing case because IPAR created to reprocess impacted claims
        due to FinThrive denials.
        These cases should be killed from workflow.
        """

        fields = extract_deterministic_ticket_extra_fields(
            text,
            _Mapper(),
            "provider_appeals_mgmt",
        )

        by_name = {field.field_name: field.new_value for field in fields}
        self.assertEqual(by_name["Category"], "Medicare Coverage Determination")
        self.assertEqual(by_name["Sub Category"], "Local Coverage Determination")
        self.assertEqual(by_name["Status"], "Overturned")
        self.assertEqual(by_name["Status Reason"], "Claims Processed Incorrectly")
        self.assertEqual(by_name["Case Closure Date"], "2026-06-04 00:00:00")
        self.assertIn("FinThrive denials", by_name["Dispute Request Reason"])
        self.assertEqual(by_name["Terminate Workflow"], "Y")

    def test_extracts_wf_close_phrase_as_workflow_termination(self) -> None:
        text = """
        CENTENE | Retail CCIU - Close Cases from WF
        TC, all cases were updated via BECU but cases are not out of WF. Please update.
        """

        fields = extract_deterministic_ticket_extra_fields(text, _Mapper(), "der")

        self.assertEqual([(field.field_name, field.new_value) for field in fields], [
            ("Terminate Workflow", "Y"),
        ])

    def test_does_not_apply_provider_aliases_to_other_modules(self) -> None:
        fields = extract_deterministic_ticket_extra_fields(
            "Category - Medicare Coverage Determination",
            _Mapper(),
            "appeals_mgmt",
        )

        self.assertEqual(fields, [])

    def test_fuzzy_labels_resolve_only_to_configured_provider_fields(self) -> None:
        text = """
        Sub catagory - Local Coverage Determination
        Status reasn - Claim Processed Incorrectly
        Case nots - Close due to corrected claim processing.
        Account Status - Active
        """

        fields = extract_deterministic_ticket_extra_fields(
            text,
            _Mapper(),
            "provider_appeals_mgmt",
        )

        by_name = {field.field_name: field.new_value for field in fields}
        self.assertEqual(by_name["Sub Category"], "Local Coverage Determination")
        self.assertEqual(by_name["Status Reason"], "Claim Processed Incorrectly")
        self.assertEqual(by_name["Dispute Request Reason"], "Close due to corrected claim processing.")
        self.assertNotIn("Status", by_name)

    def test_extracts_case_ids_from_case_list_only_spreadsheet(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["Dispute #", "Current Status"])
        ws.append(["D123456789", "In Progress"])
        ws.append(["D123456789", "In Progress"])
        ws.append(["D987654321", "In Progress"])
        out = BytesIO()
        wb.save(out)

        case_ids = extract_case_ids(out.getvalue())

        self.assertEqual(case_ids, ["D123456789", "D987654321"])
        self.assertEqual(infer_module_from_case_ids(case_ids), "provider_appeals_mgmt")

    def test_extracts_case_ids_from_csv_case_list(self) -> None:
        csv_bytes = b"der_no\n26123456789\n26123456789\n26987654321\n"

        case_ids = extract_case_ids(csv_bytes)
        module, rows = parse_spreadsheet(csv_bytes)

        self.assertEqual(case_ids, ["26123456789", "26987654321"])
        self.assertEqual(infer_module_from_case_ids(case_ids), "der")
        self.assertEqual(module, "der")
        self.assertEqual(rows, [])

    def test_extracts_case_ids_from_every_sheet(self) -> None:
        wb = Workbook()
        first = wb.active
        first.title = "Batch 1"
        first.append(["Dispute #", "Current Status"])
        first.append(["D123456789", "In Progress"])
        first.append(["D222222222", "In Progress"])

        second = wb.create_sheet("Batch 2")
        second.append(["Notes only"])
        second.append(["case D333333333 needs workflow close"])
        second.append(["D222222222"])

        third = wb.create_sheet("Batch 3")
        third.append(["Case #", "Status"])
        third.append(["G444444444", "Open"])

        out = BytesIO()
        wb.save(out)

        case_ids = extract_case_ids(out.getvalue())

        self.assertEqual(case_ids, ["D123456789", "D222222222", "D333333333", "G444444444"])

    def test_determination_value_does_not_trigger_user_termination_module(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["Case #", "Field Name", "New Value"])
        ws.append([
            "D123456789",
            "Type",
            "MD Pre-Service Medical Necessity Determination",
        ])

        out = BytesIO()
        wb.save(out)

        module, rows = parse_spreadsheet(out.getvalue())

        self.assertEqual(module, "provider_appeals_mgmt")
        self.assertEqual(rows[0].module, "provider_appeals_mgmt")
        self.assertEqual(rows[0].new_value, "MD Pre-Service Medical Necessity Determination")


if __name__ == "__main__":
    unittest.main()
