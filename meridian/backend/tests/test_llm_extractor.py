from __future__ import annotations

import json
from unittest.mock import MagicMock, patch

import pytest


# ── Helpers ──────────────────────────────────────────────────────────────────

def _make_bedrock_response(rows: list[dict]) -> dict:
    """Build a mock Bedrock response body containing JSON rows."""
    body_str = json.dumps({"rows": rows})
    return {
        "body": MagicMock(read=MagicMock(return_value=body_str.encode())),
        "ResponseMetadata": {"HTTPStatusCode": 200},
    }


def _make_mapper():
    """Return a minimal FieldMapper stub with known fields."""
    mapper = MagicMock()
    mapper._map = {
        ("appeals_mgmt", "decision date"): MagicMock(
            ui_name="Decision Date", table_name="appeals_mgmt"
        ),
        ("appeals_mgmt", "requester resolution mail date"): MagicMock(
            ui_name="Requester Resolution Mail Date", table_name="appeals_mgmt"
        ),
        ("grievance_mgmt", "status"): MagicMock(
            ui_name="Status", table_name="grievance_mgmt"
        ),
    }
    return mapper


def _make_worksheet(rows: list[list]):
    """Build a minimal openpyxl-like worksheet stub."""
    ws = MagicMock()
    ws.max_row = len(rows)
    ws.max_column = max(len(r) for r in rows) if rows else 1
    iter_rows_result = []
    for row in rows:
        cells = []
        for val in row:
            cell = MagicMock()
            cell.value = val
            cells.append(cell)
        iter_rows_result.append(tuple(cells))
    ws.iter_rows = MagicMock(return_value=iter_rows_result)
    return ws


# ── Tests ─────────────────────────────────────────────────────────────────────

class TestLLMExtractorInit:
    def test_loads_config(self):
        """LLMExtractor reads llm_config.yaml and exposes model_id and region."""
        from services.llm_extractor import LLMExtractor
        extractor = LLMExtractor()
        assert extractor.model_id
        assert extractor.region
        assert extractor.temperature == 0.0

    def test_default_model_is_haiku(self):
        from services.llm_extractor import LLMExtractor
        extractor = LLMExtractor()
        assert "haiku" in extractor.model_id.lower()


class TestBuildSystemPrompt:
    def test_prompt_contains_module_names(self):
        from services.llm_extractor import LLMExtractor
        extractor = LLMExtractor()
        mapper = _make_mapper()
        prompt = extractor.build_system_prompt(mapper)
        assert "appeals_mgmt" in prompt
        assert "grievance_mgmt" in prompt

    def test_prompt_contains_known_field_names(self):
        from services.llm_extractor import LLMExtractor
        extractor = LLMExtractor()
        mapper = _make_mapper()
        prompt = extractor.build_system_prompt(mapper)
        assert "Decision Date" in prompt
        assert "Status" in prompt

    def test_prompt_contains_anti_hallucination_rule(self):
        from services.llm_extractor import LLMExtractor
        extractor = LLMExtractor()
        mapper = _make_mapper()
        prompt = extractor.build_system_prompt(mapper)
        assert "verbatim" in prompt.lower() or "must appear" in prompt.lower()

    def test_prompt_contains_output_schema(self):
        from services.llm_extractor import LLMExtractor
        extractor = LLMExtractor()
        mapper = _make_mapper()
        prompt = extractor.build_system_prompt(mapper)
        assert "case_id" in prompt
        assert "field_name" in prompt
        assert "new_value" in prompt


class TestExtractRows:
    def test_returns_spreadsheet_rows_on_success(self):
        """Happy path: Bedrock returns token, extractor restores real case_id."""
        from services.llm_extractor import LLMExtractor
        from models.schemas import SpreadsheetRow

        ws = _make_worksheet([
            ["Case Number", "UI Section", "Old Value", "New Value"],
            ["A2601540904", "Appeals Decision Date", "3/9/2026 9:27", "2/16/2026 10:00"],
        ])
        mapper = _make_mapper()
        extractor = LLMExtractor()

        # Derive the token that _anonymise_sheet will assign for A2601540904
        _, restore_map = extractor._anonymise_sheet(ws)
        token = next(t for t, r in restore_map.items() if r == "A2601540904")

        bedrock_rows = [
            {
                "case_id": token,   # Claude returns token, not real case_id
                "module": "appeals_mgmt",
                "field_name": "Decision Date",
                "old_value": "2026-03-09 09:27:00",
                "new_value": "2026-02-16 10:00:00",
            }
        ]

        with patch.object(extractor, "_call_bedrock", return_value={"rows": bedrock_rows}):
            rows = extractor.extract(ws, mapper)

        assert len(rows) == 1
        assert isinstance(rows[0], SpreadsheetRow)
        assert rows[0].case_id == "A2601540904"   # real case_id restored
        assert rows[0].field_name == "Decision Date"
        assert rows[0].new_value == "2026-02-16 10:00:00"
        assert rows[0].module == "appeals_mgmt"

    def test_auth_number_like_case_id_is_anonymized_and_restored(self):
        """Auth numbers such as 60332RRX00 must survive LLM fallback tokenization."""
        from services.llm_extractor import LLMExtractor

        ws = _make_worksheet([
            ["Case #", "Revision #", "Field A", "Field B"],
            ["60332RRX00", "1", "Standard", "Expedited"],
        ])
        mapper = _make_mapper()
        extractor = LLMExtractor()

        _, restore_map = extractor._anonymise_sheet(ws)
        token = next(t for t, r in restore_map.items() if r == "60332RRX00")
        assert token.endswith("_H")

        bedrock_rows = [
            {
                "case_id": token,
                "module": "auth_review",
                "field_name": "Review Priority",
                "old_value": "Standard",
                "new_value": "Expedited",
            }
        ]

        with patch.object(extractor, "_call_bedrock", return_value={"rows": bedrock_rows}):
            rows = extractor.extract(ws, mapper)

        assert len(rows) == 1
        assert rows[0].case_id == "60332RRX00"
        assert rows[0].module == "auth_review"
        assert rows[0].field_name == "Review Priority"

    def test_returns_empty_on_bedrock_error(self):
        """If Bedrock raises any exception, extract() returns [] without crashing."""
        from services.llm_extractor import LLMExtractor

        ws = _make_worksheet([["some", "random", "headers"]])
        mapper = _make_mapper()
        extractor = LLMExtractor()

        with patch.object(extractor, "_call_bedrock", side_effect=Exception("network error")):
            rows = extractor.extract(ws, mapper)

        assert rows == []

    def test_returns_empty_on_none_from_bedrock(self):
        """If _call_bedrock returns None, extract() returns []."""
        from services.llm_extractor import LLMExtractor

        ws = _make_worksheet([["something"]])
        mapper = _make_mapper()
        extractor = LLMExtractor()

        with patch.object(extractor, "_call_bedrock", return_value=None):
            rows = extractor.extract(ws, mapper)

        assert rows == []

    def test_skips_rows_with_missing_case_id(self):
        """Rows without case_id are dropped — anti-hallucination."""
        from services.llm_extractor import LLMExtractor

        ws = _make_worksheet([["A2601540904", "Decision Date", "", "2/16/2026"]])
        mapper = _make_mapper()
        extractor = LLMExtractor()
        _, restore_map = extractor._anonymise_sheet(ws)
        token = next(t for t, r in restore_map.items() if r == "A2601540904")

        bedrock_rows = [
            {"case_id": "", "module": "appeals_mgmt",
             "field_name": "Decision Date", "new_value": "2026-02-16 10:00:00"},
            {"case_id": token, "module": "appeals_mgmt",
             "field_name": "Decision Date", "new_value": "2026-02-16 10:00:00"},
        ]

        with patch.object(extractor, "_call_bedrock", return_value={"rows": bedrock_rows}):
            rows = extractor.extract(ws, mapper)

        assert len(rows) == 1
        assert rows[0].case_id == "A2601540904"

    def test_skips_rows_with_missing_field_name(self):
        """Rows without field_name are dropped."""
        from services.llm_extractor import LLMExtractor

        bedrock_rows = [
            {"case_id": "A2601540904", "module": "appeals_mgmt",
             "field_name": "", "new_value": "2026-02-16 10:00:00"},
        ]
        ws = _make_worksheet([["A2601540904", "", "", "2/16/2026"]])
        mapper = _make_mapper()
        extractor = LLMExtractor()

        with patch.object(extractor, "_call_bedrock", return_value={"rows": bedrock_rows}):
            rows = extractor.extract(ws, mapper)

        assert rows == []

    def test_sets_llm_extracted_reason(self):
        """SpreadsheetRow returned by LLM extractor has reason='llm_extracted'."""
        from services.llm_extractor import LLMExtractor

        ws = _make_worksheet([["G26062567435", "Status", "Resolved", "Void"]])
        mapper = _make_mapper()
        extractor = LLMExtractor()
        _, restore_map = extractor._anonymise_sheet(ws)
        token = next(t for t, r in restore_map.items() if r == "G26062567435")

        bedrock_rows = [
            {"case_id": token, "module": "grievance_mgmt",
             "field_name": "Status", "old_value": "Resolved", "new_value": "Void"},
        ]

        with patch.object(extractor, "_call_bedrock", return_value={"rows": bedrock_rows}):
            rows = extractor.extract(ws, mapper)

        assert rows[0].reason == "llm_extracted"

    def test_drops_hallucinated_token(self):
        """Rows with a token not in restore_map are dropped (hallucination guard)."""
        from services.llm_extractor import LLMExtractor

        ws = _make_worksheet([["A2601540904", "Decision Date", "", "2/16/2026"]])
        mapper = _make_mapper()
        extractor = LLMExtractor()

        # Claude returns a token that wasn't in the sheet
        bedrock_rows = [
            {"case_id": "CASE_999_X", "module": "appeals_mgmt",
             "field_name": "Decision Date", "new_value": "2026-02-16 10:00:00"},
        ]

        with patch.object(extractor, "_call_bedrock", return_value={"rows": bedrock_rows}):
            rows = extractor.extract(ws, mapper)

        assert rows == []


class TestCallBedrock:
    def test_calls_invoke_model(self):
        """_call_bedrock calls invoke_model on the boto3 client with correct payload."""
        from services.llm_extractor import LLMExtractor
        import json

        extractor = LLMExtractor()
        response_data = {"rows": [{"case_id": "A123", "module": "appeals_mgmt",
                                    "field_name": "Decision Date", "new_value": "2026-01-01 10:00:00"}]}
        mock_body = MagicMock()
        mock_body.read.return_value = json.dumps({
            "content": [{"text": json.dumps(response_data)}]
        }).encode()
        mock_client = MagicMock()
        mock_client.invoke_model.return_value = {"body": mock_body}

        with patch.object(extractor, "_bedrock_client", mock_client):
            result = extractor._call_bedrock("system prompt", "user message")

        assert mock_client.invoke_model.called
        call_kwargs = mock_client.invoke_model.call_args[1]
        body = json.loads(call_kwargs["body"])
        assert body["messages"][0]["role"] == "user"
        assert body["messages"][0]["content"] == "user message"
        assert body["system"] == "system prompt"
        assert body["temperature"] == 0.0
        assert result == response_data


class TestPhiRedaction:
    def test_redacts_common_phi_before_llm(self):
        from services.llm_extractor import _contains_residual_phi_risk, _redact_sensitive_text

        text = (
            "First Name: John; Member John Smith lives at 123 Main Street. "
            "Email john.smith@example.com phone 555-123-4567 SSN 123-45-6789."
        )

        redacted = _redact_sensitive_text(text)

        assert "John Smith" not in redacted
        assert "123 Main Street" not in redacted
        assert "john.smith@example.com" not in redacted
        assert "555-123-4567" not in redacted
        assert "123-45-6789" not in redacted
        assert _contains_residual_phi_risk(redacted) is False

    def test_phi_label_terms_are_loaded_from_config(self):
        from services.llm_extractor import _load_phi_label_terms, _redact_sensitive_text

        terms = _load_phi_label_terms()

        assert "first name" in terms
        assert "subscriber id" in terms
        assert "provider phone" in terms
        assert "Jane" not in _redact_sensitive_text("first name: Jane")
        assert "ABC123" not in _redact_sensitive_text("subscriber_id: ABC123")

    def test_residual_phi_risk_detector_fails_closed(self):
        from services.llm_extractor import _contains_residual_phi_risk

        assert _contains_residual_phi_risk("Unredacted member appears as Smith, Jane") is True


class TestSpreadsheetParserFallback:
    def test_parser_calls_llm_when_no_rows(self):
        """parse_spreadsheet falls back to LLM when standard parse returns empty."""
        import io
        import openpyxl
        from unittest.mock import patch, MagicMock
        from services.spreadsheet_parser import parse_spreadsheet
        from models.schemas import SpreadsheetRow

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Case Number", "UI Section", "Old Value", "New Value", "Comments"])
        ws.append(["A2601540904", "Appeals Decision Date", "3/9/2026", "2/16/2026 10:00", "Audit"])
        buf = io.BytesIO()
        wb.save(buf)
        file_bytes = buf.getvalue()

        mock_rows = [SpreadsheetRow(
            row_num=1, module="appeals_mgmt", case_id="A2601540904",
            field_name="Decision Date", new_value="2026-02-16 10:00:00",
            reason="llm_extracted",
        )]

        with patch("services.spreadsheet_parser.LLMExtractor") as MockExtractor:
            instance = MockExtractor.return_value
            instance.extract.return_value = mock_rows
            with patch("services.spreadsheet_parser._parse_tall", return_value=(None, [])):
                with patch("services.spreadsheet_parser._parse_wide", return_value=(None, [])):
                    module, rows = parse_spreadsheet(file_bytes)

        assert len(rows) == 1
        assert rows[0].case_id == "A2601540904"
        assert rows[0].reason == "llm_extracted"

    def test_parser_calls_llm_when_deterministic_rows_are_low_confidence(self):
        """Rows like Field A/Field B are parsed output, but not usable Meridian fields."""
        import io
        import openpyxl
        from unittest.mock import patch
        from services.field_mapper import FieldMapper, FieldMapping
        from services.spreadsheet_parser import parse_spreadsheet
        from models.schemas import SpreadsheetRow

        mapper = FieldMapper()
        mapping = FieldMapping(
            table_name="auth_review",
            ui_name="Review Priority",
            column_name="auth_priority_id",
            ref_table=None,
            ref_id_col=None,
            ref_desc_col=None,
            lov_expression=None,
        )
        mapper._map[(mapping.table_name, mapping.ui_name.lower())] = mapping
        mapper._by_ui.setdefault(mapping.ui_name.lower(), []).append(mapping)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([
            "Case #",
            "Revision #",
            "Field A (Old Value)",
            "Field A (NEW Value)",
            "Field B (Old Value)",
            "Field B (NEW Value)",
        ])
        ws.append([
            "Case #",
            "Revision #",
            "Review Priority (Old)",
            "Review Priority (NEW)",
            "Due Date (Old)",
            "Due Date (NEW)",
        ])
        ws.append([
            "60332RRX00",
            1,
            "Standard",
            "Expedited",
            "2026-05-25 23:59:00",
            "2026-05-21 11:13:00",
        ])
        buf = io.BytesIO()
        wb.save(buf)

        mock_rows = [
            SpreadsheetRow(
                row_num=1,
                module="auth_review",
                case_id="60332RRX00",
                revision=1,
                field_name="Review Priority",
                old_value="Standard",
                new_value="Expedited",
                reason="llm_extracted",
            )
        ]

        with patch("services.spreadsheet_parser.LLMExtractor") as MockExtractor:
            instance = MockExtractor.return_value
            instance.extract.return_value = mock_rows
            module, rows = parse_spreadsheet(buf.getvalue(), mapper=mapper)

        MockExtractor.assert_called_once()
        instance.extract.assert_called_once()
        assert module == "auth_review"
        assert len(rows) == 1
        assert rows[0].case_id == "60332RRX00"
        assert rows[0].field_name == "Review Priority"
        assert rows[0].reason == "llm_extracted"

    def test_parser_does_not_call_llm_when_standard_parse_succeeds(self):
        """LLM is NOT called when standard parser finds rows."""
        import io
        import openpyxl
        from unittest.mock import patch
        from services.spreadsheet_parser import parse_spreadsheet

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Case #", "Field Name", "Old Value", "New Value"])
        ws.append(["A2601540904", "Decision Date", "2026-03-09", "2026-02-16"])
        buf = io.BytesIO()
        wb.save(buf)
        file_bytes = buf.getvalue()

        with patch("services.spreadsheet_parser.LLMExtractor") as MockExtractor:
            parse_spreadsheet(file_bytes)
            MockExtractor.assert_not_called()
